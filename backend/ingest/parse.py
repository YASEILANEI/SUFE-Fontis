"""解析 规章制度/ 下的官方文档为纯文本，输出元数据索引。

产出：
  backend/data/parsed/<id>.txt      每个文档一个文本文件
  backend/data/documents.json       元数据索引 [{id, title, category, rel_path, format, pages, note}]
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]          # CampusAgent/
DOCS_DIR = ROOT / "规章制度"
DATA_DIR = ROOT / "backend" / "data"
PARSED_DIR = DATA_DIR / "parsed"

SUPPORTED = {".pdf", ".docx", ".xlsx", ".xls"}
UNSUPPORTED = {".doc"}  # 老格式，需转换后才能解析


def parse_pdf(path: Path) -> tuple[str, dict]:
    import fitz  # pymupdf

    doc = fitz.open(str(path))
    pages = []
    for page in doc:
        pages.append(page.get_text())
    text = "\n".join(pages)
    return text, {"pages": doc.page_count}


def parse_docx(path: Path) -> tuple[str, dict]:
    import docx

    d = docx.Document(str(path))
    parts: list[str] = []
    for p in d.paragraphs:
        if p.text.strip():
            parts.append(p.text.strip())
    for tbl in d.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts), {}


def parse_spreadsheet(path: Path) -> tuple[str, dict]:
    import pandas as pd

    sheets = pd.read_excel(str(path), sheet_name=None, header=None)
    lines: list[str] = []
    for name, df in sheets.items():
        lines.append(f"[工作表：{name}]")
        for _, row in df.iterrows():
            vals = [str(v) for v in row.tolist() if pd.notna(v)]
            if vals:
                lines.append(" | ".join(vals))
    return "\n".join(lines), {"sheets": len(sheets)}


def parse_file(path: Path) -> dict | None:
    ext = path.suffix.lower()
    if ext == ".pdf":
        text, meta = parse_pdf(path)
    elif ext == ".docx":
        text, meta = parse_docx(path)
    elif ext in (".xls", ".xlsx"):
        text, meta = parse_spreadsheet(path)
    elif ext in UNSUPPORTED:
        return None
    else:
        return None

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()

    rel_path = path.relative_to(ROOT).as_posix()
    category = path.parent.relative_to(DOCS_DIR).parts[0] if path.is_relative_to(DOCS_DIR) else "其他"
    note = "scan" if ext == ".pdf" and len(text.strip()) < 20 else ""

    return {
        "id": f"{category}/{path.stem}",
        "title": re.sub(r"^[0-9a-fA-F]{32}", "", path.stem).strip(),
        "category": category,
        "rel_path": rel_path,
        "format": ext.lstrip("."),
        "note": note,
        "text": text,
        **meta,
    }


def main() -> list[dict]:
    PARSED_DIR.mkdir(parents=True, exist_ok=True)
    docs: list[dict] = []
    skipped: list[str] = []
    for path in sorted(DOCS_DIR.rglob("*")):
        if not path.is_file():
            continue
        parsed = parse_file(path)
        if parsed is None:
            skipped.append(path.as_posix())
            continue
        docs.append(parsed)

    # 索引只保留元数据，text 单独按 id 存文件（id 含 /，转成安全文件名）
    index = [{k: v for k, v in d.items() if k != "text"} for d in docs]
    (DATA_DIR / "documents.json").write_text(
        json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    for d in docs:
        safe = d["id"].replace("/", "__")
        (PARSED_DIR / f"{safe}.txt").write_text(d["text"], encoding="utf-8")

    print(f"解析完成：{len(docs)} 个文档")
    for d in docs:
        flag = " [扫描件]" if d["note"] == "scan" else ""
        print(f"  [{d['category']}] {d['title']}{flag} ({len(d['text'])} 字)")
    if skipped:
        print(f"跳过 {len(skipped)} 个文件：")
        for s in skipped:
            print(f"  {s}")
    return docs


if __name__ == "__main__":
    main()
